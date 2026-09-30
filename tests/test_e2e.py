"""Сквозной цикл: обезличить → проверить → «ответ Claude» → вернуть."""
import os
import re

import docx
import openpyxl
import pytest

from shirma import morph
from shirma.cli import main
from shirma.formats import texts
from shirma.registry import Registry
from shirma.workspace import Workspace, walk_files

from make_samples import make_all

SECRETS = ['Иванов', 'Ивана', 'Сеитов', 'Айгерим', 'Омаров', 'Кузнецов', 'Ромашк', 'Ақжол', 'Бәйтерек',
           'Сәрсенбаев', 'Абдыкаримов', '850315300128', '920708450152', '051212650159', '51212650159',
           '050140004877', 'KZ86125KZT5004100100', '4400 4301 2345 6789', '123-45-67', '1234567', '555 12 34',
           'romashka', 'akzhol', 'ivanov', 'СОКОЛ', 'Сокол']


@pytest.fixture(scope='module')
def ws(tmp_path_factory):
    root = tmp_path_factory.mktemp('ws')
    assert main(['init', str(root)]) == 0
    w = Workspace(str(root))
    make_all(w.orig)
    assert main(['obezlichit', '--root', str(root)]) == 0
    return w


def _all_text(w, folder):
    out = []
    for rel in walk_files(folder):
        with open(os.path.join(folder, rel), 'rb') as f:
            out += [rel] + texts(f.read(), rel)
    return '\n'.join(out)


def test_copies_have_no_secrets(ws):
    blob = _all_text(ws, ws.copies)
    for s in SECRETS:
        # само значение или его падежная форма (но не часть другого слова: «Сокол» ≠ «Сокольникова»)
        m = re.search(rf'(?<![^\W\d_]){re.escape(s)}[^\W\d_]{{0,3}}(?![^\W\d_])', blob)
        assert not m, (s, blob[max(0, m.start() - 40):m.end() + 40])


def test_check_passes(ws):
    assert main(['proverit', '--root', ws.root]) == 0


def test_formatting_and_numbers_kept(ws):
    d = docx.Document(os.path.join(ws.copies, [f for f in walk_files(ws.copies) if f.endswith('.docx')][0]))
    assert d.paragraphs[0].runs[0].bold
    assert 'Сумма договора: 12 500 000 тенге, дата 01.02.2025.' in [p.text for p in d.paragraphs]
    wb = openpyxl.load_workbook(os.path.join(ws.copies, 'Реестр.xlsx'))
    sh = wb['Сотрудники']
    assert sh['E4'].value == 150000 and sh['E5'].value == 80000000000
    assert sh['G4'].value == '=A4&" — "&C4'
    assert wb.properties.creator == 'shirma'


def test_fakes_are_valid(ws):
    from shirma.fakes import kz_valid, luhn_valid, iban_valid
    reg = Registry(ws.db)
    for t in ('iin', 'bin'):
        for fake in reg.nums[t].values():
            assert kz_valid(fake)
    for fake in reg.nums['card'].values():
        assert luhn_valid(fake)
    for fake in reg.nums['iban'].values():
        assert iban_valid(fake)


def _fake_person(reg, sur):
    for p in reg.persons:
        if p.orig['sur'].startswith(sur):
            return p
    raise AssertionError(sur)


def test_restore_claude_output(ws):
    reg = Registry(ws.db)
    iv = _fake_person(reg, 'Иванов').fake
    ku = _fake_person(reg, 'Кузнецов').fake
    org = reg.orgs['ромашка'].fake['name']
    phone = reg.nums['phone']['7011234567']
    g = 'm'
    # падежи, которых не было во входящих
    iv_ablt = morph.decline_surname(iv['sur'], g)[4] + ' ' + morph.decline_name(iv['name'], g)[4]
    ku_datv = morph.decline_surname(ku['sur'], 'f')[2] + ' ' + morph.decline_name(ku['name'], 'f')[2] + ' ' + \
        morph.decline_patr(ku['patr'], 'f')[2]
    text = (f'Встреча с {iv_ablt} и ТОО «{org}». Письмо — {ku_datv}. '
            f'Звонить: 8-{phone[:3]}-{phone[3:6]}-{phone[6:8]}-{phone[8:]}. '
            f'Почта {reg.emails["ivanov@romashka.kz"]}.')
    os.makedirs(ws.claude_out, exist_ok=True)
    d = docx.Document()
    d.add_paragraph(text)
    d.save(os.path.join(ws.claude_out, f'Итоги {iv["sur"]}.docx'))
    with open(os.path.join(ws.claude_out, 'итоги.md'), 'w', encoding='utf-8') as f:
        f.write('# ' + text + '\n')
    wb = openpyxl.Workbook()
    wb.active.append(['ФИО', 'ИИН'])
    wb.active.append([iv['sur'] + ' ' + iv['name'] + ' ' + iv['patr'], int(reg.nums['iin']['850315300128'])])
    wb.save(os.path.join(ws.claude_out, 'свод.xlsx'))
    with open(os.path.join(ws.claude_out, 'отчёт.pdf'), 'wb') as f:
        f.write(b'%PDF-1.4 fake')

    assert main(['vernut', '--root', ws.root]) == 0
    res = os.listdir(ws.result)
    assert 'Итоги Иванов.docx' in res
    d = docx.Document(os.path.join(ws.result, 'Итоги Иванов.docx'))
    t = d.paragraphs[0].text
    assert 'Ивановым Иваном' in t
    assert 'ТОО «Ромашка»' in t
    assert 'Кузнецовой Марии Петровне' in t
    assert '8-701-123-45-67' in t
    assert 'ivanov@romashka.kz' in t
    with open(os.path.join(ws.result, 'итоги.md'), encoding='utf-8') as f:
        assert f.read().startswith('# Встреча с Ивановым Иваном')
    wb = openpyxl.load_workbook(os.path.join(ws.result, 'свод.xlsx'))
    assert wb.active['A2'].value == 'Иванов Иван Иванович'
    assert wb.active['B2'].value == 850315300128
    assert 'отчёт.pdf' in res


def test_line_breaks_inside_paragraph(ws):
    """Перенос строки и табуляция в абзаце docx — граница слова, а не склейка кусков."""
    import io
    from shirma import formats
    from shirma.replace import FWD, Stats
    reg = Registry(ws.db)
    d = docx.Document()
    d.add_paragraph('БИН 050140004877\nИванову И. И.\nivanov@romashka.kz\tКузнецова М.П.')
    buf = io.BytesIO()
    d.save(buf)
    out = formats.convert(buf.getvalue(), 'письмо.docx', reg, FWD, Stats(), [])
    fwd = '\n'.join(texts(out, 'письмо.docx'))
    for s in ('Иванов', 'Кузнецов', 'ivanov', '050140004877'):
        assert s not in fwd

    iv = _fake_person(reg, 'Иванов').fake
    ku = _fake_person(reg, 'Кузнецов').fake
    iv_datv = morph.decline_surname(iv['sur'], 'm')[2]
    d = docx.Document()
    d.add_paragraph(f'БИН {reg.nums["bin"]["050140004877"]}\n{iv_datv} {iv["name"][0]}. {iv["patr"][0]}.\n'
                    f'{reg.emails["ivanov@romashka.kz"]}\t{ku["sur"]} {ku["name"][0]}.{ku["patr"][0]}.')
    d.save(os.path.join(ws.claude_out, 'письмо.docx'))
    assert main(['vernut', '--root', ws.root]) == 0
    t = docx.Document(os.path.join(ws.result, 'письмо.docx')).paragraphs[0].text
    assert 'Иванову И. И.' in t
    assert 'ivanov@romashka.kz' in t
    assert 'Кузнецова М.П.' in t
    assert '050140004877' in t


def test_second_run_is_idempotent(ws):
    before = {f: os.path.getmtime(os.path.join(ws.copies, f)) for f in walk_files(ws.copies)}
    assert main(['obezlichit', '--root', ws.root]) == 0
    after = {f: os.path.getmtime(os.path.join(ws.copies, f)) for f in walk_files(ws.copies)}
    assert before == after


def test_old_file_redone_when_new_person_learned(ws):
    # старый файл упоминает человека без инициалов — его не узнать; новый файл даёт ФИО полностью
    with open(os.path.join(ws.orig, 'старое.txt'), 'w', encoding='utf-8') as f:
        f.write('Звонил Тлеулесов, просил перезвонить.\n')
    assert main(['obezlichit', '--root', ws.root]) == 0
    with open(os.path.join(ws.orig, 'новое.txt'), 'w', encoding='utf-8') as f:
        f.write('Директор: Тлеулесов Арман Серикович.\n')
    assert main(['obezlichit', '--root', ws.root]) == 0
    blob = _all_text(ws, ws.copies)
    assert 'Тлеулесов' not in blob
