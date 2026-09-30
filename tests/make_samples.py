"""Синтетические документы для тестов (все данные выдуманы)."""
import os

import docx
import openpyxl
from docx.shared import Pt


def make_docx(path):
    d = docx.Document()
    d.core_properties.author = 'Иванов Иван Иванович'
    sec = d.sections[0]
    sec.header.paragraphs[0].text = 'ТОО «Ромашка», БИН 050140004877'
    sec.footer.paragraphs[0].text = 'Исполнитель: Сеитова А.Б., тел. +7 (701) 123-45-67'
    p = d.add_paragraph()
    p.add_run('Договор между ТОО «Ромаш').bold = True
    p.add_run('ка» в лице директора Ив')
    r = p.add_run('анова Ивана Ивановича')
    r.italic = True
    r.font.size = Pt(14)
    p.add_run(' и ИП Сәрсенбаев Н.Қ.')
    d.add_paragraph('ИИН директора 850315300128, счёт KZ86125KZT5004100100, почта ivanov@romashka.kz, '
                    'сайт https://www.romashka.kz/about')
    d.add_paragraph('Согласовано с Айгерим Сеитовой. Копию направить И.И. Иванову.')
    t = d.add_table(rows=2, cols=3)
    t.rows[0].cells[0].text = 'ФИО'
    t.rows[0].cells[1].text = 'ИИН'
    t.rows[0].cells[2].text = 'Телефон'
    t.rows[1].cells[0].text = 'Омарова Жанна Серікқызы'
    t.rows[1].cells[1].text = '920708450152'
    t.rows[1].cells[2].text = '8 777 555 12 34'
    d.add_paragraph('Сумма договора: 12 500 000 тенге, дата 01.02.2025.')
    d.save(path)


def make_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Сотрудники'
    ws.append(['Реестр сотрудников'])
    ws.append([])
    ws.append(['ФИО', 'ИИН', 'Контрагент', 'Телефон', 'Сумма', 'Комментарий'])
    ws.append(['Иванов Иван Иванович', 850315300128, 'Ай', '+77011234567', 150000, 'звонил Сеитовой'])
    ws.append(['Кузнецова М.П.', 51212650159, 'ТОО «Ромашка»', 87775551234, 80000000000, ''])
    ws.append(['СОКОЛ ПЁТР', '051212650159', 'Бәйтерек', '8 (7172) 55-44-33', 1, ''])
    ws['G4'] = '=A4&" — "&C4'
    ws2 = wb.create_sheet('Иванов')
    ws2['A1'] = 'Отчёт Иванова И.И.'
    ws2['B1'] = "=Сотрудники!A4"
    wb.properties.creator = 'Иванов И.И.'
    wb.save(path)


def make_txt(path):
    with open(path, 'w', encoding='cp1251') as f:
        f.write('Служебная записка\nОт: Иванов И.И.\nКому: Кузнецовой Марии Петровне\n'
                'Прошу оплатить счёт ТОО «Ромашка» (БИН 050140004877) на карту 4400 4301 2345 6789.\n')


def make_csv(path):
    with open(path, 'w', encoding='utf-8-sig', newline='') as f:
        f.write('Контрагент;БИН;Контакт;Email\r\n'
                'ТОО «Ақжол»;050140004877;Нурлан Абдыкаримов;n.abdykarimov@akzhol.kz\r\n'
                'Ай;;Сокол П.;info@ai.kz\r\n')


def make_all(folder):
    os.makedirs(folder, exist_ok=True)
    make_docx(os.path.join(folder, 'Договор Иванов И.И..docx'))
    make_xlsx(os.path.join(folder, 'Реестр.xlsx'))
    make_txt(os.path.join(folder, 'Записка.txt'))
    make_csv(os.path.join(folder, 'Контрагенты.csv'))


if __name__ == '__main__':
    import sys
    make_all(sys.argv[1])
