"""Рабочая папка, в пути к которой есть русские буквы и пробелы."""
import json
import os
import subprocess
import unicodedata

import docx
import pytest

from shirma.cli import main
from shirma.hook import allowed_commands, check
from shirma.registry import Registry
from shirma.setup import _posix, _python
from shirma.workspace import Workspace, walk_files

from make_samples import make_all


@pytest.fixture(scope='module')
def ws(tmp_path_factory):
    root = tmp_path_factory.mktemp('base') / 'Документы Ивановой' / 'Работа с Claude (ТОО «Ромашка»)'
    root.mkdir(parents=True)
    assert main(['init', str(root)]) == 0
    w = Workspace(str(root))
    make_all(w.private_in)
    return w


def test_full_cycle(ws):
    assert main(['obfuscate', '--root', ws.root]) == 0
    assert main(['check', '--root', ws.root]) == 0
    copies = walk_files(ws.claude_in)
    assert len(copies) == 4 and not any('Иванов' in c for c in copies)

    reg = Registry(ws.db)
    iv = next(p for p in reg.persons if p.orig['sur'] == 'Иванов').fake
    d = docx.Document()
    d.add_paragraph(f'Итог: {iv["sur"]} {iv["name"]} {iv["patr"]}')
    d.save(os.path.join(ws.claude_out, f'Итог {iv["sur"]}.docx'))
    assert main(['restore', '--root', ws.root]) == 0
    t = docx.Document(os.path.join(ws.private_out, 'Итог Иванов.docx')).paragraphs[0].text
    assert t == 'Итог: Иванов Иван Иванович'


def test_workspace_found_from_claude_folder(ws, monkeypatch):
    monkeypatch.chdir(os.path.join(ws.claude, 'input'))
    found = Workspace.find()
    assert found and found.root == ws.root
    assert main(['check']) == 0


def test_settings_and_launchers(ws):
    with open(os.path.join(ws.claude, '.claude', 'settings.json'), encoding='utf-8') as f:
        settings = json.load(f)
    assert any('Документы Ивановой' in r for r in settings['permissions']['deny'])
    hook_cmd = settings['hooks']['PreToolUse'][0]['hooks'][0]['command']
    assert 'Документы Ивановой' in hook_cmd
    if os.name != 'nt':
        for title in ('1-Обезличить', '2-Проверить', '3-Вернуть'):
            p = os.path.join(ws.root, f'{title}.command')
            assert os.access(p, os.X_OK)
            subprocess.run(['bash', '-n', p], check=True)


def test_allowed_command_runs(ws):
    """Команда, разрешённая Claude, действительно работает с русским путём."""
    cmd = next(c for c in allowed_commands(_posix(_python()), _posix(ws.root)) if ' check ' in c)
    if os.name == 'nt':
        pytest.skip('проверяется на POSIX-оболочке')
    r = subprocess.run(['bash', '-c', cmd], capture_output=True, text=True, cwd=ws.claude)
    assert r.returncode == 0, r.stdout + r.stderr
    assert 'Проверено файлов' in r.stdout


def _hook(ws, tool, inp):
    return check({'tool_name': tool, 'tool_input': inp, 'cwd': ws.claude},
                 ws.private, ws.root, _posix(_python()))


def test_hook_with_cyrillic_root(ws):
    assert _hook(ws, 'Read', {'file_path': os.path.join(ws.private, 'system', 'mapping.csv')})
    assert _hook(ws, 'Read', {'file_path': '../private/input/Записка.txt'})
    assert _hook(ws, 'Bash', {'command': f'ls "{ws.root}"'})
    assert _hook(ws, 'Read', {'file_path': 'input/Записка.txt'}) is None
    assert _hook(ws, 'Bash', {'command': f'ls "{ws.claude}/input"'}) is None
    cmd = next(c for c in allowed_commands(_posix(_python()), _posix(ws.root)) if ' restore ' in c)
    assert _hook(ws, 'Bash', {'command': cmd}) is None


def test_hook_nfd_path(ws):
    """macOS может отдавать путь в NFD (й = и + ˘) — блокировка не должна зависеть от формы."""
    nfd = unicodedata.normalize('NFD', os.path.join(ws.private, 'input', 'Записка.txt'))
    assert nfd != unicodedata.normalize('NFC', nfd)
    assert _hook(ws, 'Read', {'file_path': nfd})
    assert _hook(ws, 'Bash', {'command': f'cat "{nfd}"'})


def test_hook_process_with_cyrillic_args(ws):
    """Хук как отдельный процесс (так его вызывает Claude Code) с русскими путями в аргументах."""
    with open(os.path.join(ws.claude, '.claude', 'settings.json'), encoding='utf-8') as f:
        hook_cmd = json.load(f)['hooks']['PreToolUse'][0]['hooks'][0]['command']
    if os.name == 'nt':
        pytest.skip('проверяется на POSIX-оболочке')
    bad = json.dumps({'tool_name': 'Read', 'tool_input': {'file_path': '../private/system/mapping.csv'},
                      'cwd': ws.claude}, ensure_ascii=False)
    r = subprocess.run(['bash', '-c', hook_cmd], input=bad, capture_output=True, text=True, cwd=ws.claude)
    assert r.returncode == 2 and 'private' in r.stderr
    good = json.dumps({'tool_name': 'Read', 'tool_input': {'file_path': 'input/x.docx'}, 'cwd': ws.claude})
    r = subprocess.run(['bash', '-c', hook_cmd], input=good, capture_output=True, text=True, cwd=ws.claude)
    assert r.returncode == 0, r.stderr


def test_skill_created(ws):
    p = os.path.join(ws.claude, '.claude', 'skills', 'shirma', 'SKILL.md')
    with open(p, encoding='utf-8') as f:
        text = f.read()
    assert text.startswith('---\nname: shirma\ndescription: ')
    for c in allowed_commands(_posix(_python()), _posix(ws.root)):
        assert c in text          # точные команды, разрешённые в настройках
    assert '{cmd_' not in text
    with open(os.path.join(ws.claude, '.claude', 'settings.json'), encoding='utf-8') as f:
        deny = json.load(f)['permissions']['deny']
    assert 'Edit(./.claude/**)' in deny and 'Write(./CLAUDE.md)' in deny
